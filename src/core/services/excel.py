from django.utils.timezone import localtime
import openpyxl

from core.models import CatequeseInfantilModel, CrismaModel, Perseveranca_MEJ_Model, CatequeseAdultoModel


def gerar_Workbook():
    qs = CatequeseInfantilModel.objects.all().order_by('id')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Catequese Infantil"
    list_header = ['Nome','Sexo','Data de Nascimento','Naturalidade','Nome do Pai','Nome da Mãe','Endereço','Cidade','UF','Celular do Pai','Celular da Mãe',]
    list_header += ['Batizado','Data do Batismo','Diocese do Batismo','Paróquia do Batismo','Celebrante do Batismo']
    list_header += ['Horário da Catequese']
    list_header += ['Possui Deficiência', 'Descrição da Deficiência', 'Possui Transtorno', 'Descrição do Transtorno']
    list_header += ['Medicamento de Uso Contínuo',  'Descrição do Medicamento', 'Horário do Medicamento']
    list_header += ['Acompanhamento Psicológico','Descrição do Acompanhamento',]
    list_header += ['Nome do Responsável','CPF do Responsável','Endereço do Responsável']
    list_header += ['Ficha Impressa','Ficha Assinada','Data de Registro']
    ws.append(list_header)
    for registro in qs:
        ws.append([
            registro.nome,
            registro.sexo,
            registro.data_nascimento.strftime('%d/%m/%Y') if registro.data_nascimento else '',
            registro.naturalidade,
            registro.nome_pai,
            registro.nome_mae,
            registro.endereco,
            registro.cidade,
            registro.uf,
            registro.celular_pai,
            registro.celular_mae,
            'Sim' if registro.batizado else 'Não',
            registro.batizado_data.strftime('%d/%m/%Y') if registro.batizado_data else 'NÃO BATIZADO',
            registro.batizado_diocese,
            registro.batizado_paroquia,
            registro.batizado_celebrante,
            registro.turma.nome,
            'Sim' if registro.possui_deficiencia else 'Não',
            registro.descricao_deficiencia,
            'Sim' if registro.possui_transtorno else 'Não',
            registro.descricao_transtorno,
            'Sim' if registro.medicamento_uso_continuo else 'Não',
            registro.descricao_medicamento,
            registro.medicamento_horario,
            'Sim' if registro.acompanhamento_psicologico else 'Não',
            registro.descricao_acompanhamento,
            registro.nome_responsavel,
            registro.cpf_responsavel,
            registro.endereco_responsavel,
            "Sim" if registro.ficha_impressa else "Não",
            "Sim" if registro.ficha_assinada else "Não",    
            localtime(registro.criado_em).strftime('%d/%m/%Y - %H:%M:%S') if registro.criado_em else '',
        ])
    ws = wb.create_sheet(title="Crisma")
    
    qs = CrismaModel.objects.all().order_by('id')
    list_header = [
        'Nome',
        'Sexo',
        'Data de Nascimento',
        'Naturalidade',
        'Nome do Pai',
        'Nome da Mãe',
        'Endereço',
        'Cidade',
        'UF',
        'Celular do Pai',
        'Celular da Mãe',
]

    list_header += [
        'Batizado',
        'Data do Batismo',
        'Diocese do Batismo',
        'Paróquia do Batismo',
        'Celebrante do Batismo',
    ]

    list_header += [
        'Primeira Eucaristia',
        'Data da Primeira Eucaristia',
        'Diocese da Primeira Eucaristia',
        'Paróquia da Primeira Eucaristia',
        'Celebrante da Primeira Eucaristia',
    ]

    list_header += [
        'Horário da Crisma',
        'Nome do Padrinho',
        'Celular do Padrinho',
    ]

    list_header += [
        'Possui Deficiência',
        'Descrição da Deficiência',
        'Possui Transtorno',
        'Descrição do Transtorno',
    ]

    list_header += [
        'Medicamento de Uso Contínuo',
        'Descrição do Medicamento',
        'Horário do Medicamento',
    ]

    list_header += [
        'Acompanhamento Psicológico',
        'Descrição do Acompanhamento',
    ]

    list_header += [
        'Nome do Responsável',
        'CPF do Responsável',
        'Endereço do Responsável',
    ]

    list_header += [
        'Ficha Impressa',
        'Ficha Assinada',
        'Data de Registro',
    ]

    ws.append(list_header)

    for registro in qs:
        ws.append([
            registro.nome,
            registro.sexo,
            registro.data_nascimento.strftime('%d/%m/%Y') if registro.data_nascimento else '',
            registro.naturalidade,
            registro.nome_pai,
            registro.nome_mae,
            registro.endereco,
            registro.cidade,
            registro.uf,
            registro.celular_pai,
            registro.celular_mae,

            'Sim' if registro.batizado else 'Não',
            registro.batizado_data.strftime('%d/%m/%Y') if registro.batizado_data else 'NÃO BATIZADO',
            registro.batizado_diocese,
            registro.batizado_paroquia,
            registro.batizado_celebrante,

            'Sim' if registro.primeira_eucaristia else 'Não',
            registro.primeira_eucaristia_data.strftime('%d/%m/%Y') if registro.primeira_eucaristia_data else 'NÃO REALIZADA',
            registro.primeira_eucaristia_diocese,
            registro.primeira_eucaristia_paroquia,
            registro.primeira_eucaristia_celebrante,

            registro.turma.nome,
            registro.padrinho_nome,
            registro.padrinho_celular,

            'Sim' if registro.possui_deficiencia else 'Não',
            registro.descricao_deficiencia,

            'Sim' if registro.possui_transtorno else 'Não',
            registro.descricao_transtorno,

            'Sim' if registro.medicamento_uso_continuo else 'Não',
            registro.descricao_medicamento,
            registro.medicamento_horario,

            'Sim' if registro.acompanhamento_psicologico else 'Não',
            registro.descricao_acompanhamento,

            registro.nome_responsavel,
            registro.cpf_responsavel,
            registro.endereco_responsavel,

            'Sim' if registro.ficha_impressa else 'Não',
            'Sim' if registro.ficha_assinada else 'Não',

            localtime(registro.criado_em).strftime('%d/%m/%Y - %H:%M:%S') if registro.criado_em else '',
        ])
    
    ws = wb.create_sheet(title="Perseverança MEJ")
    
    qs = Perseveranca_MEJ_Model.objects.all().order_by('id')
    list_header = [
        'Nome',
        'Sexo',
        'Data de Nascimento',
        'Naturalidade',
        'Nome do Pai',
        'Nome da Mãe',
        'Endereço',
        'Cidade',
        'UF',
        'Celular do Pai',
        'Celular da Mãe',
    ]

    list_header += [
        'Batizado',
        'Data do Batismo',
        'Diocese do Batismo',
        'Paróquia do Batismo',
        'Celebrante do Batismo',
    ]

    list_header += [
        'Primeira Eucaristia',
        'Data da Primeira Eucaristia',
        'Diocese da Primeira Eucaristia',
        'Paróquia da Primeira Eucaristia',
        'Celebrante da Primeira Eucaristia',
    ]

    list_header += [
        'Horário da Perseverança / MEJ',
    ]

    list_header += [
        'Possui Deficiência',
        'Descrição da Deficiência',
        'Possui Transtorno',
        'Descrição do Transtorno',
    ]

    list_header += [
        'Medicamento de Uso Contínuo',
        'Descrição do Medicamento',
        'Horário do Medicamento',
    ]

    list_header += [
        'Acompanhamento Psicológico',
        'Descrição do Acompanhamento',
    ]

    list_header += [
        'Nome do Responsável',
        'CPF do Responsável',
        'Endereço do Responsável',
    ]

    list_header += [
        'Ficha Impressa',
        'Ficha Assinada',
        'Data de Registro',
    ]

    ws.append(list_header)

    for registro in qs:
        ws.append([
            registro.nome,
            registro.sexo,
            registro.data_nascimento.strftime('%d/%m/%Y') if registro.data_nascimento else '',
            registro.naturalidade,
            registro.nome_pai,
            registro.nome_mae,
            registro.endereco,
            registro.cidade,
            registro.uf,
            registro.celular_pai,
            registro.celular_mae,

            'Sim' if registro.batizado else 'Não',
            registro.batizado_data.strftime('%d/%m/%Y') if registro.batizado_data else 'NÃO BATIZADO',
            registro.batizado_diocese,
            registro.batizado_paroquia,
            registro.batizado_celebrante,

            'Sim' if registro.primeira_eucaristia else 'Não',
            registro.primeira_eucaristia_data.strftime('%d/%m/%Y') if registro.primeira_eucaristia_data else 'NÃO REALIZADA',
            registro.primeira_eucaristia_diocese,
            registro.primeira_eucaristia_paroquia,
            registro.primeira_eucaristia_celebrante,

            registro.get_horario_display(),

            'Sim' if registro.possui_deficiencia else 'Não',
            registro.descricao_deficiencia,

            'Sim' if registro.possui_transtorno else 'Não',
            registro.descricao_transtorno,

            'Sim' if registro.medicamento_uso_continuo else 'Não',
            registro.descricao_medicamento,
            registro.medicamento_horario,

            'Sim' if registro.acompanhamento_psicologico else 'Não',
            registro.descricao_acompanhamento,

            registro.nome_responsavel,
            registro.cpf_responsavel,
            registro.endereco_responsavel,

            'Sim' if registro.ficha_impressa else 'Não',
            'Sim' if registro.ficha_assinada else 'Não',

            localtime(registro.criado_em).strftime('%d/%m/%Y - %H:%M:%S') if registro.criado_em else '',
        ])
    ws = wb.create_sheet(title="Catequese Adulto")
    qs = CatequeseAdultoModel.objects.all().order_by('id')
    
    list_header = [
        'Nome',
        'Sexo',
        'Data de Nascimento',
        'Naturalidade',
        'CPF',
        'Celular',
        'Nome do Pai',
        'Nome da Mãe',
        'Endereço',
        'Cidade',
        'UF',
        'Estado Civil',
    ]

    list_header += [
        'Batizado',
        'Data do Batismo',
        'Diocese do Batismo',
        'Paróquia do Batismo',
        'Celebrante do Batismo',
    ]

    list_header += [
        'Primeira Eucaristia',
        'Data da Primeira Eucaristia',
        'Diocese da Primeira Eucaristia',
        'Paróquia da Primeira Eucaristia',
        'Celebrante da Primeira Eucaristia',
    ]

    list_header += [
        'Casado na Igreja',
        'Data do Casamento',
        'Diocese do Casamento',
        'Paróquia do Casamento',
        'Celebrante do Casamento',
    ]

    list_header += [
        'Horário da Catequese (Adulto)',
        'Nome do Padrinho',
        'Celular do Padrinho',
    ]

    list_header += [
        'Ficha Impressa',
        'Ficha Assinada',
        'Data de Registro',
    ]

    ws.append(list_header)

    for registro in qs:
        ws.append([
            registro.nome,
            registro.sexo,
            registro.data_nascimento.strftime('%d/%m/%Y') if registro.data_nascimento else '',
            registro.naturalidade,
            registro.cpf,
            registro.celular,
            registro.nome_pai,
            registro.nome_mae,
            registro.endereco,
            registro.cidade,
            registro.uf,
            registro.estado_civil,

            'Sim' if registro.batizado else 'Não',
            registro.batizado_data.strftime('%d/%m/%Y') if registro.batizado_data else 'NÃO BATIZADO',
            registro.batizado_diocese,
            registro.batizado_paroquia,
            registro.batizado_celebrante,

            'Sim' if registro.primeira_eucaristia else 'Não',
            registro.primeira_eucaristia_data.strftime('%d/%m/%Y') if registro.primeira_eucaristia_data else 'NÃO REALIZADA',
            registro.primeira_eucaristia_diocese,
            registro.primeira_eucaristia_paroquia,
            registro.primeira_eucaristia_celebrante,

            'Sim' if registro.casado_igreja else 'Não',
            registro.casado_igreja_data.strftime('%d/%m/%Y') if registro.casado_igreja_data else 'NÃO INFORMADO',
            registro.casado_igreja_diocese,
            registro.casado_igreja_paroquia,
            registro.casado_igreja_celebrante,

            registro.get_horario_display(),
            registro.padrinho_nome,
            registro.padrinho_celular,

            'Sim' if registro.ficha_impressa else 'Não',
            'Sim' if registro.ficha_assinada else 'Não',

            localtime(registro.criado_em).strftime('%d/%m/%Y - %H:%M:%S') if registro.criado_em else '',
        ])
    return wb